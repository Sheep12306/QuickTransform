"""Excel workbook to CSV conversion."""

import csv
import zipfile
from pathlib import Path

from openpyxl import load_workbook
import xlrd


def convert_excel_to_csvs(source, output_dir, preview_limit=200):
    source = Path(source)
    output_dir = Path(output_dir)
    if not source.is_file():
        raise FileNotFoundError("找不到待转换的 Excel 文件")

    output_dir.mkdir(parents=True, exist_ok=True)
    sheets = _read_sheets(source)
    for index, sheet in enumerate(sheets):
        csv_path = output_dir / "{0}.csv".format(_safe_sheet_name(sheet["name"], index))
        _write_csv(sheet["rows"], csv_path)
        sheet["csv_path"] = csv_path
        sheet["total_rows"] = len(sheet["rows"])
        sheet["columns"] = max((len(row) for row in sheet["rows"]), default=0)
        sheet["preview"] = sheet["rows"][:preview_limit]

    zip_path = output_dir / "all-sheets.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for sheet in sheets:
            archive.write(sheet["csv_path"], sheet["csv_path"].name)

    return sheets, zip_path


def _read_sheets(source):
    with open(source, "rb") as file:
        head = file.read(8)
    if head[:4] == b"\xd0\xcf\x11\xe0":
        return _read_xls_sheets(source)
    return _read_xlsx_sheets(source)


def _read_xlsx_sheets(source):
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            rows = _trim_rows(list(worksheet.iter_rows(values_only=True)))
            sheets.append({"name": worksheet.title, "rows": rows})
        return sheets
    finally:
        workbook.close()


def _read_xls_sheets(source):
    workbook = xlrd.open_workbook(str(source))
    sheets = []
    for worksheet in workbook.sheets():
        rows = []
        for row_index in range(worksheet.nrows):
            rows.append([worksheet.cell_value(row_index, col_index) for col_index in range(worksheet.ncols)])
        sheets.append({"name": worksheet.name, "rows": _trim_rows(rows)})
    return sheets


def _write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file, lineterminator="\r\n")
        writer.writerows(rows)


def _trim_rows(rows):
    cleaned = [list(row) for row in rows]
    while cleaned and all(cell is None or str(cell) == "" for cell in cleaned[-1]):
        cleaned.pop()
    return cleaned


def _safe_sheet_name(name, index):
    cleaned = "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in str(name))
    cleaned = cleaned.strip("._")
    return cleaned or "sheet-{0}".format(index + 1)
