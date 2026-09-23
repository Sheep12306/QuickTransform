"""PDF to Word conversion using pdf2docx locally."""

from pathlib import Path

from docx import Document
from pdf2docx import Converter


def convert_pdf_to_word(source, destination):
    source = Path(source)
    destination = Path(destination)
    if not source.is_file():
        raise FileNotFoundError("找不到待转换的 PDF 文件")

    converter = Converter(str(source))
    try:
        converter.convert(
            str(destination),
            start=0,
            end=None,
            clip_image_res_ratio=6.0,
        )
    finally:
        converter.close()

    _normalize_docx_runs(destination)

    if not destination.is_file():
        raise RuntimeError("PDF 转 Word 失败：没有生成 Word 文件")


def _normalize_docx_runs(path):
    document = Document(path)
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    for paragraph in paragraphs:
        _normalize_paragraph_runs(paragraph)

    document.save(path)


def _normalize_paragraph_runs(paragraph):
    merged = []
    for run in list(paragraph.runs):
        _apply_font_fallbacks(run)
        key = _run_key(run)
        if merged and merged[-1][0] == key:
            previous = merged[-1][1]
            previous.text = previous.text + run.text
            run._r.getparent().remove(run._r)
        else:
            merged.append((key, run))


def _apply_font_fallbacks(run):
    name = (run.font.name or "").lower()
    if run.bold is not True and ("bold" in name or "black" in name or "heavy" in name):
        run.bold = True
    if run.italic is not True and ("italic" in name or "oblique" in name):
        run.italic = True


def _run_key(run):
    font = run.font
    color = None
    if font.color and font.color.rgb is not None:
        color = str(font.color.rgb)
    return (
        font.name,
        font.size.pt if font.size else None,
        run.bold,
        run.italic,
        font.underline,
        color,
        run.style.style_id if run.style else None,
    )
